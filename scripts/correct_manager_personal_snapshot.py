#!/usr/bin/env python3
"""Repair manager rows that were incorrectly populated with store KPI values."""

import argparse
import json
from calendar import monthrange
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


PERSON_METRICS = {
    "AQ": "TTL AQ上線點數",
    "A999": "AQ V+D 999 (含)以上",
    "A1399": "AQ V+D 1399 (含)以上",
    "RT": "RT上線點數",
    "R999": "RT V+D 999 (含)以上",
    "R1399": "RT V+D 1399 (含)以上",
    "好速": "好速案銷售點數",
    "特維": "特殊維繫用戶續約數",
    "配件": "配件及其他營收",
    "包膜": "包膜與保貼營收",
}


def manager_rows(path: Path) -> dict[str, dict]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    sheet = workbook["上線數KPI_個人達成率_明細"]
    result = {}
    for row in sheet.iter_rows(min_row=10, max_col=10, values_only=True):
        role = str(row[4] or "")
        if not str(row[2] or "").startswith("DNB") or role not in ("店長", "代理店長"):
            continue
        result[str(row[6] or "")] = {
            "employee_id": str(row[5] or ""),
            "rank": int(row[8]),
            "overall_rate": float(row[9]),
        }
    workbook.close()
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dashboard", type=Path, required=True)
    parser.add_argument("--kpicalc", type=Path, required=True)
    parser.add_argument("--current-xlsx", type=Path, required=True)
    parser.add_argument("--previous-xlsx", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    dashboard = json.loads(args.dashboard.read_text(encoding="utf-8"))
    kpicalc = json.loads(args.kpicalc.read_text(encoding="utf-8"))
    current = manager_rows(args.current_xlsx)
    previous = manager_rows(args.previous_xlsx)
    calc_people = {str(row.get("pname") or ""): row for row in kpicalc.get("persons", [])}
    cutoff = date.fromisoformat(dashboard["kpiBattle"]["data_as_of_date"])
    elapsed = cutoff.day / monthrange(cutoff.year, cutoff.month)[1]

    repaired = 0
    for row in dashboard["kpiBattle"].get("personal", []):
        if str(row.get("role") or "") not in ("店長", "代理店長"):
            continue
        name = str(row.get("name") or "")
        source = calc_people.get(name)
        identity = current.get(name)
        if not source or not identity:
            raise RuntimeError(f"missing manager source: {name}")
        prior = previous.get(name)
        row["rank"] = identity["rank"]
        row["overall_rate"] = float(source["official"])
        row["overall_rate_dod"] = (
            round(identity["overall_rate"] - prior["overall_rate"], 4) if prior else None
        )
        row["rank_dod"] = prior["rank"] - identity["rank"] if prior else None
        metrics = {}
        for label, source_key in PERSON_METRICS.items():
            metric = source.get("items", {}).get(source_key)
            if metric is None:
                raise RuntimeError(f"missing {source_key}: {name}")
            actual = metric.get("a")
            target = metric.get("t")
            daily_target = round(float(target) * elapsed, 1) if target is not None else None
            metrics[label] = {
                "actual": actual,
                "target": target,
                "daily_target": daily_target,
                "daily_gap": round(float(actual) - daily_target, 1) if actual is not None and daily_target is not None else None,
                "rate": metric.get("reportRate"),
                "dod": None,
            }
        row["metrics"] = metrics
        repaired += 1

    if repaired != len(current) or repaired != 9:
        raise RuntimeError(f"manager count mismatch: repaired={repaired}, source={len(current)}")
    dashboard["kpiBattle"]["personal_semantics"] = "individual-v1"
    dashboard["kpiBattle"]["personal_source_sheet"] = "上線數KPI_個人達成率_明細"
    args.output.write_text(json.dumps(dashboard, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


if __name__ == "__main__":
    main()
