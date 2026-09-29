#!/usr/bin/env python3
"""Render the verified `好速上線明細` worksheet to complete PNG page(s).

This renderer is deliberately driven by worksheet rows rather than a fixed
capture range.  It also writes a JSON receipt containing every row actually
drawn so the mail publisher can compare the image payload with the workbook
before publishing.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


ERROR_CODE = "GOODSPEED_PNG_ROW_MISMATCH"
SHEET_NAME = "好速上線明細"
HEADERS = (
    "店點", "統計日期", "同仁", "認列狀態", "產品／方案", "頻寬", "約期",
    "36M加碼", "500M", "1G", "速率加碼", "當日公司認列點數", "當日戰報含另加點數",
)
KEY_COLUMNS = (0, 1, 2, 4, 11, 12)
DEFAULT_ROWS_PER_PAGE = 12


def fail(message: str) -> "None":
    raise RuntimeError(f"{ERROR_CODE}: {message}")


def display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:g}"
    if isinstance(value, Decimal):
        return format(value, "f").rstrip("0").rstrip(".")
    return str(value).strip()


def number(value: str) -> Decimal:
    text = str(value or "0").replace(",", "").strip()
    try:
        return Decimal(text)
    except Exception as exc:
        fail(f"invalid numeric value {value!r}: {exc}")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass(frozen=True)
class Row:
    values: tuple[str, ...]

    @property
    def key(self) -> str:
        return "|".join(self.values[index] for index in KEY_COLUMNS)


def read_rows(workbook_path: Path) -> tuple[str, list[Row], list[str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        fail(f"worksheet {SHEET_NAME!r} is missing")
    worksheet = workbook[SHEET_NAME]
    title = "北一二B 好速上線明細"
    header_row = None
    footers: list[str] = []
    rows: list[Row] = []
    for raw in worksheet.iter_rows(values_only=True):
        values = tuple(display(value) for value in raw[: len(HEADERS)])
        if values[0].startswith("北一二B 好速上線明細"):
            title = values[0]
            continue
        if values[0] == HEADERS[0] and values[1] == HEADERS[1]:
            header_row = True
            continue
        if not header_row:
            continue
        nonempty = sum(bool(value) for value in values)
        if nonempty == 0:
            continue
        if values[0].startswith(("資料日：", "戰報加碼（")):
            footers.append(values[0])
            continue
        if nonempty < 7:
            fail(f"incomplete worksheet row after header: {values!r}")
        row = Row(values)
        if not all(row.values[index] for index in KEY_COLUMNS[:5]):
            fail(f"stable key is incomplete: {row.values!r}")
        rows.append(row)
    if header_row is None:
        fail("header row was not found")
    return title, rows, footers


def find_font(explicit: str | None, bold: bool = False) -> str:
    candidates = [explicit] if explicit else []
    candidates.extend([
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/Library/Fonts/NotoSansCJKtc-Regular.otf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold
        else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ])
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    fail("no usable font was found; set --font to a Traditional Chinese font")


def fit_text(draw: ImageDraw.ImageDraw, value: str, font_path: str, max_width: int,
             start_size: int = 27, minimum: int = 16) -> ImageFont.FreeTypeFont:
    for size in range(start_size, minimum - 1, -1):
        font = ImageFont.truetype(font_path, size)
        if draw.textbbox((0, 0), value or " ", font=font)[2] <= max_width:
            return font
    return ImageFont.truetype(font_path, minimum)


def page_paths(output: Path, page_count: int) -> list[Path]:
    if page_count == 1:
        return [output]
    return [output.with_name(f"{output.stem}_p{index}{output.suffix}")
            for index in range(1, page_count + 1)]


def render(workbook_path: Path, output: Path, receipt_path: Path,
           rows_per_page: int, font_path: str | None = None) -> dict:
    if rows_per_page < 1:
        fail("rows_per_page must be positive")
    title, rows, footers = read_rows(workbook_path)
    page_count = max(1, math.ceil(len(rows) / rows_per_page))
    outputs = page_paths(output, page_count)
    regular_path = find_font(font_path, False)
    bold_path = find_font(font_path, True)
    widths = [250, 150, 145, 175, 175, 115, 100, 120, 105, 105, 130, 155, 165]
    margin, title_h, header_h, row_h, footer_h = 16, 78, 82, 62, 112
    width = sum(widths) + margin * 2
    drawn_rows: list[dict] = []
    files: list[dict] = []
    for page_index, path in enumerate(outputs):
        start = page_index * rows_per_page
        page_rows = rows[start : start + rows_per_page]
        height = margin * 2 + title_h + header_h + row_h * len(page_rows) + footer_h
        image = Image.new("RGB", (width, height), "#ffffff")
        draw = ImageDraw.Draw(image)
        title_font = ImageFont.truetype(bold_path, 40)
        draw.rectangle((0, 0, width, title_h + margin), fill="#173361")
        page_label = f"（{page_index + 1}/{page_count}）" if page_count > 1 else ""
        title_text = f"{title}{page_label}"
        bbox = draw.textbbox((0, 0), title_text, font=title_font)
        draw.text(((width - (bbox[2] - bbox[0])) / 2, 20), title_text,
                  font=title_font, fill="#ffffff")
        y = title_h + margin
        x = margin
        for index, (header, cell_w) in enumerate(zip(HEADERS, widths)):
            draw.rectangle((x, y, x + cell_w, y + header_h), fill="#ffcb18", outline="#c8d2df", width=2)
            label = header.replace("當日公司認列點數", "當日公司\n認列點數").replace(
                "當日戰報含另加點數", "當日戰報\n含另加點數")
            font = ImageFont.truetype(bold_path, 23)
            lines = label.split("\n")
            line_h = 27
            top = y + (header_h - line_h * len(lines)) / 2
            for line_no, line in enumerate(lines):
                box = draw.textbbox((0, 0), line, font=font)
                draw.text((x + (cell_w - (box[2] - box[0])) / 2, top + line_no * line_h),
                          line, font=font, fill="#173361")
            x += cell_w
        for local_index, row in enumerate(page_rows):
            y += header_h if local_index == 0 else row_h
            x = margin
            drawn_cells = []
            for column, (value, cell_w) in enumerate(zip(row.values, widths)):
                fill = "#f4f7fb" if (start + local_index) % 2 == 0 else "#ffffff"
                if column == 3:
                    fill = "#d7f5df"
                draw.rectangle((x, y, x + cell_w, y + row_h), fill=fill, outline="#c8d2df", width=2)
                font = fit_text(draw, value, regular_path, cell_w - 12)
                box = draw.textbbox((0, 0), value or " ", font=font)
                draw.text((x + (cell_w - (box[2] - box[0])) / 2,
                           y + (row_h - (box[3] - box[1])) / 2 - box[1]),
                          value, font=font, fill="#142033")
                if value:
                    drawn_cells.append(column)
                x += cell_w
            if len(drawn_cells) < 7:
                fail(f"rendered row has too few populated cells: {row.values!r}")
            drawn_rows.append({"key": row.key, "values": list(row.values),
                               "page": page_index + 1, "row": local_index + 1,
                               "drawn_cells": drawn_cells})
        table_bottom = title_h + margin + header_h + row_h * len(page_rows)
        footer_font = ImageFont.truetype(regular_path, 22)
        receipt = f"逐列核對：{len(rows)}/{len(rows)} 筆｜公司點數 {sum(number(r.values[11]) for r in rows)}｜戰報點數 {sum(number(r.values[12]) for r in rows)}"
        footer_texts = ([receipt] + footers[:2]) if page_index == page_count - 1 else [receipt]
        for line_no, line in enumerate(footer_texts):
            draw.text((margin + 8, table_bottom + 18 + line_no * 29), line,
                      font=footer_font, fill="#173361")
        path.parent.mkdir(parents=True, exist_ok=True)
        image.save(path, format="PNG", optimize=True)
        files.append({"path": str(path), "bytes": path.stat().st_size,
                      "sha256": sha256(path), "rows": len(page_rows)})
    expected = [row.key for row in rows]
    actual = [row["key"] for row in drawn_rows]
    if expected != actual:
        fail(f"expected row keys {expected!r}, rendered {actual!r}")
    receipt = {
        "schema_version": "north12b-goodspeed-png-receipt/v1",
        "status": "pass",
        "source": str(workbook_path),
        "source_sha256": sha256(workbook_path),
        "row_count": len(rows),
        "company_points": str(sum(number(row.values[11]) for row in rows)),
        "battle_points": str(sum(number(row.values[12]) for row in rows)),
        "row_keys": expected,
        "drawn_rows": drawn_rows,
        "files": files,
    }
    receipt_path.parent.mkdir(parents=True, exist_ok=True)
    receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return receipt


def verify(workbook_path: Path, receipt_path: Path) -> dict:
    _, rows, _ = read_rows(workbook_path)
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    expected = [row.key for row in rows]
    if receipt.get("status") != "pass" or receipt.get("row_keys") != expected:
        fail("receipt row keys do not match workbook")
    if receipt.get("row_count") != len(rows):
        fail("receipt row count does not match workbook")
    if receipt.get("company_points") != str(sum(number(row.values[11]) for row in rows)):
        fail("company point total does not match workbook")
    if receipt.get("battle_points") != str(sum(number(row.values[12]) for row in rows)):
        fail("battle point total does not match workbook")
    for item in receipt.get("files", []):
        path = Path(item["path"])
        if not path.exists() or path.stat().st_size == 0 or sha256(path) != item.get("sha256"):
            fail(f"rendered PNG does not match receipt: {path}")
    return receipt


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    make = sub.add_parser("render")
    make.add_argument("workbook", type=Path)
    make.add_argument("output", type=Path)
    make.add_argument("--receipt", type=Path)
    make.add_argument("--rows-per-page", type=int, default=DEFAULT_ROWS_PER_PAGE)
    make.add_argument("--font")
    check = sub.add_parser("verify")
    check.add_argument("workbook", type=Path)
    check.add_argument("receipt", type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        if args.command == "render":
            receipt_path = args.receipt or args.output.with_suffix(".receipt.json")
            result = render(args.workbook, args.output, receipt_path, args.rows_per_page, args.font)
        else:
            result = verify(args.workbook, args.receipt)
        print(json.dumps({"status": "pass", "row_count": result["row_count"],
                          "files": result["files"]}, ensure_ascii=False))
        return 0
    except Exception as exc:
        message = str(exc)
        if not message.startswith(ERROR_CODE):
            message = f"{ERROR_CODE}: {message}"
        print(message)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
