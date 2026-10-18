"""Personal QIS: source-backed reported-score scenarios, never inferred raw totals."""
from decimal import Decimal, InvalidOperation, ROUND_CEILING
from pathlib import Path
from datetime import date
from collections import Counter
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import hashlib, re, unicodedata
from openpyxl import load_workbook

def text(value):return unicodedata.normalize('NFKC',str('' if value is None else value)).strip()

def number(value):
    if value is None or isinstance(value,bool) or text(value)=='':return None
    try:
        n=Decimal(str(value))
        return n if n.is_finite() else None
    except InvalidOperation:return None

def estimate_100_cases(count,score):
    n,s=number(count),number(score)
    result={'status':'pending','case_count':None,'current_score':None,'additional_100':None,'projected_report_average':None}
    if n is None or n<0 or n!=n.to_integral_value():return result
    result['case_count']=int(n)
    if n==0:return {**result,'status':'no_data','additional_100':1,'projected_report_average':100.0}
    if s is None or s<0 or s>100:return result
    result['current_score']=float(s)
    if s>=98:return {**result,'status':'at_least_98','additional_100':0,'projected_report_average':float(s)}
    k=int((n*(Decimal(98)-s)/2).to_integral_value(rounding=ROUND_CEILING))
    # This is a reported-average scenario, not the employee's original total score.
    simulated=(n*s+100*k)/(n+k)
    prior=(n*s+100*(k-1))/(n+k-1)
    if simulated<98 or prior>=98:raise ValueError('QIS_100_CASE_SCENARIO_RECONCILIATION_FAILED')
    return {**result,'status':'below_98','additional_100':k,'projected_report_average':float(simulated)}

def choose_column(hits,label):
    if not hits:raise ValueError('QIS_PERSONAL_IDENTITY_COLUMN_MISSING:'+label)
    best=max(hits.values());cols=[col for col,n in hits.items() if n==best]
    if len(cols)!=1:raise ValueError('QIS_PERSONAL_IDENTITY_COLUMN_AMBIGUOUS:'+label)
    return cols[0]

def raw_cells(path,sheet):
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(path) as z:
        wb=ET.fromstring(z.read('xl/workbook.xml'))
        rid=next(x.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
            for x in wb.find('m:sheets',ns) if x.attrib['name']==sheet)
        rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        target=next(x.attrib['Target'] for x in rels if x.attrib['Id']==rid)
        member=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        tree=ET.fromstring(z.read(member));out={}
        for cell in tree.findall('.//m:c',ns):
            v=cell.find('m:v',ns);f=cell.find('m:f',ns)
            numeric=cell.attrib.get('t') in (None,'n')
            out[cell.attrib['r']]={'raw':v.text if v is not None and numeric else None,'has_formula':f is not None,
                'kind':'formula-cached' if f is not None else 'numeric-xml-value' if v is not None and numeric
                    else 'non-numeric-xml-value' if v is not None else 'missing'}
        return out

def extract_personal_qis(path,roster,report_date,scope='北一二B'):
    path=Path(path);wb=load_workbook(path,data_only=True,read_only=True)
    if 'QIS個績' not in wb.sheetnames:raise ValueError('QIS_PERSONAL_SHEET_MISSING')
    ws=wb['QIS個績'];values=list(ws.iter_rows(values_only=True));wb.close()
    headers=[]
    for rn,row in enumerate(values[:30],1):
        counts=[i for i,x in enumerate(row) if text(x)=='計件數']
        scores=[i for i,x in enumerate(row) if text(x)=='原始分數']
        if counts:
            if len(counts)!=1 or len(scores)!=1:raise ValueError('QIS_PERSONAL_HEADER_MISMATCH')
            headers.append((rn,counts[0],scores[0]))
    if len(headers)!=1:raise ValueError('QIS_PERSONAL_HEADER_MISMATCH')
    header,count_col,score_col=headers[0]
    periods=[]
    for row in values[:header]:
        for value in row:
            match=re.fullmatch(r'(\d{4})/(\d{1,2})/(\d{1,2})\s*[~～]\s*(?:(\d{4})/)?(\d{1,2})/(\d{1,2})',text(value))
            if match:
                periods.append(date(int(match[4] or match[1]),int(match[5]),int(match[6])).isoformat())
    if len(set(periods))!=1 or periods[0]>report_date:raise ValueError('QIS_PERSONAL_CUTOFF_MISSING_OR_INVALID')
    source=[(rn,row) for rn,row in enumerate(values,1) if rn>header and text(scope) in map(text,row)]
    people={text(p['employee_id']):p for p in roster['persons']}
    stores={text(s['code']):s['name'] for s in roster['stores']}
    employee_hits,store_hits,name_hits=Counter(),Counter(),Counter()
    for _,row in source:
        matched=[people[text(v)] for v in row if text(v) in people]
        for i,v in enumerate(row):
            if text(v) in people:employee_hits[i]+=1
            if text(v) in stores:store_hits[i]+=1
            if any(text(v)==text(p['pname']) for p in matched):name_hits[i]+=1
    employee_col=choose_column(employee_hits,'employee');store_col=choose_column(store_hits,'store-code');name_col=choose_column(name_hits,'name')
    if len({employee_col,store_col,name_col,count_col,score_col})!=5:raise ValueError('QIS_PERSONAL_IDENTITY_COLUMN_OVERLAP')
    raw=raw_cells(path,'QIS個績');rows=[];seen=set()
    from openpyxl.utils import get_column_letter
    for rn,row in source:
        emp,name,store_code=text(row[employee_col]),text(row[name_col]),text(row[store_col])
        if not emp and not name:continue # formatted/summary rows are not employees
        if emp and emp in seen:raise ValueError('QIS_PERSONAL_DUPLICATE_EMPLOYEE')
        seen.add(emp)
        score_cell=f'{get_column_letter(score_col+1)}{rn}';count_cell=f'{get_column_letter(count_col+1)}{rn}'
        score_info=raw.get(score_cell,{})
        estimate=estimate_100_cases(raw.get(count_cell,{}).get('raw'),score_info.get('raw'))
        if not emp or not name or store_code not in stores:estimate['status']='pending';estimate['additional_100']=None
        store=stores.get(store_code,'待核店點');store='台北三創' if '三創' in store else store
        rows.append({**estimate,'store':store,'name':name or '待核姓名','employee_id':emp,'source_row':rn,
            'score_cell':score_cell,'count_cell':count_cell,'source_score_raw':score_info.get('raw'),
            'source_count_raw':raw.get(count_cell,{}).get('raw'),'score_value_kind':score_info.get('kind','missing')})
    below=sorted([r for r in rows if r['status']=='below_98'],key=lambda r:(-r['additional_100'],r['current_score'],r['source_row']))
    return {'schema':'north12b-personal-qis/v1','sourceSha256':hashlib.sha256(path.read_bytes()).hexdigest(),
        'cutoffDate':periods[0],'scope':scope,'header':{'row':header,'countColumn':count_col+1,'scoreColumn':score_col+1},
        'identityColumns':{'employee':employee_col+1,'storeCode':store_col+1,'name':name_col+1},'rows':rows,'below98':below,
        'noData':[r for r in rows if r['status']=='no_data'],'pending':[r for r in rows if r['status']=='pending'],
        'precisionNote':'依報表分數等權試算；臨界可能差1筆，實際以公司更新為準；未反推原始總分。'}


def verify_personal_qis(personal, source_values, source_sha256, report_date, personal_cutoff, worksheet):
    """Reject missing, duplicated or changed employee targets before attachment sealing."""
    assert personal['sourceSha256']==source_sha256,'QIS_PERSONAL_SOURCE_MISMATCH'
    assert personal['cutoffDate']==personal_cutoff<=report_date,'QIS_PERSONAL_DATE_MISMATCH'
    qsource=source_values;header=personal['header'];identity=personal['identityColumns']
    assert qsource[header['row']-1][header['countColumn']-1]=='計件數'
    assert qsource[header['row']-1][header['scoreColumn']-1]=='原始分數'
    scoped=[(rn,row) for rn,row in enumerate(qsource,1) if rn>header['row'] and '北一二B' in row
        and (row[identity['employee']-1] or row[identity['name']-1])]
    assert len(scoped)==len(personal['rows']),'QIS_PERSONAL_ROSTER_MISMATCH'
    for person,(rn,source) in zip(personal['rows'],scoped):
        assert person['source_row']==rn and person['employee_id']==str(source[identity['employee']-1])
        assert person['name']==(source[identity['name']-1] or '待核姓名')
        n=person['case_count'];score=person['current_score'];status=person['status']
        if n is not None:assert n==source[header['countColumn']-1],'QIS_PERSONAL_COUNT_MISMATCH'
        if score is not None:assert score==source[header['scoreColumn']-1],'QIS_PERSONAL_SCORE_MISMATCH'
        if n is not None:assert Decimal(person['source_count_raw'])==Decimal(str(source[header['countColumn']-1])),'QIS_PERSONAL_RAW_COUNT_MISMATCH'
        if score is not None:assert Decimal(person['source_score_raw'])==Decimal(str(source[header['scoreColumn']-1])),'QIS_PERSONAL_RAW_SCORE_MISMATCH'
        if status=='no_data':
            assert n==0 and score is None and person['additional_100']==1
        elif status=='below_98':
            assert n>0 and score<98 and score==source[header['scoreColumn']-1]
            # Independently check the threshold inequalities using exact original XML decimal text.
            raw_score=Decimal(person['source_score_raw']);count=Decimal(person['source_count_raw']);extra=person['additional_100']
            assert extra>0 and (count*raw_score+100*extra)/(count+extra)>=98
            assert (count*raw_score+100*(extra-1))/(count+extra-1)<98
        elif status=='at_least_98':assert n>0 and score>=98 and person['additional_100']==0
        else:assert status=='pending' and person['additional_100'] is None
    assert personal['below98']==sorted([p for p in personal['rows'] if p['status']=='below_98'],
        key=lambda p:(-p['additional_100'],p['current_score'],p['source_row'])), 'QIS_PERSONAL_FOLLOWUP_COVERAGE_MISMATCH'
    for key,status in [('noData','no_data'),('pending','pending')]:
        assert personal[key]==[p for p in personal['rows'] if p['status']==status],'QIS_PERSONAL_STATE_COVERAGE_MISMATCH'
    row=18
    for person in personal['below98']:
        assert [worksheet.cell(row,col).value for col in range(1,6)]==[
            person['store'],person['name'],person['current_score'],person['case_count'],person['additional_100']], 'QIS_PERSONAL_DISPLAY_MISMATCH'
        row+=1
    if not personal['below98']:row+=1
    for key in ['noData','pending']:
        if not personal[key]:continue
        row+=3
        for person in personal[key]:
            score='尚無資料' if key=='noData' else person['current_score'] if person['current_score'] is not None else '待核'
            needed='首筆100即100' if key=='noData' else '待核'
            assert [worksheet.cell(row,col).value for col in range(1,6)]==[
                person['store'],person['name'],score,person['case_count'] if person['case_count'] is not None else '待核',needed], 'QIS_PERSONAL_STATE_DISPLAY_MISMATCH'
            row+=1
    assert personal['precisionNote'] in [r[0].value for r in worksheet.iter_rows()], 'QIS_PERSONAL_PRECISION_NOTE_MISSING'

def verify_personal_qis_render(personal, receipt, png_path, workbook_path, renderer_path, font_dir):
    """Bind every displayed followup to the actual PNG, workbook, renderer and fonts."""
    from PIL import Image
    digest=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
    assert receipt['status']=='pass' and receipt['dimensions']==[1920,1080], 'QIS_PERSONAL_CANVAS_MISMATCH'
    assert receipt['allTextFitsAtSelectedFontSize'] and receipt['fontShrinking'] is False, 'QIS_PERSONAL_TEXT_OVERFLOW'
    with Image.open(png_path) as image:assert image.size==(1920,1080), 'QIS_PERSONAL_PNG_SIZE_MISMATCH'
    assert receipt['pngSha256']==digest(png_path), 'QIS_PERSONAL_PNG_HASH_MISMATCH'
    assert receipt['workbookSha256']==digest(workbook_path), 'QIS_PERSONAL_WORKBOOK_HASH_MISMATCH'
    assert receipt['rendererSha256']==digest(renderer_path), 'QIS_PERSONAL_RENDERER_HASH_MISMATCH'
    assert receipt['fontSha256']=={n:digest(Path(font_dir)/n) for n in ['NotoSansTC-Regular.ttf','NotoSansTC-Bold.ttf']}, 'QIS_PERSONAL_FONT_HASH_MISMATCH'
    assert receipt['personalCutoff']==personal['cutoffDate'].replace('-','/'), 'QIS_PERSONAL_RENDER_DATE_MISMATCH'
    people=personal['below98']+personal['noData']+personal['pending']
    drawn=[r for r in receipt['drawnRows'] if r['kind']=='people']
    assert len(drawn)==receipt['personalRows']==len(people), 'QIS_PERSONAL_RENDER_COVERAGE_MISMATCH'
    assert receipt['below98']==len(personal['below98']) and receipt['noData']==len(personal['noData']), 'QIS_PERSONAL_RENDER_STATE_MISMATCH'
    for rendered,person in zip(drawn,people):
        status=person['status']
        score='尚無資料' if status=='no_data' else person['current_score'] if person['current_score'] is not None else '待核'
        needed='首筆100即100' if status=='no_data' else '待核' if status=='pending' else person['additional_100']
        assert rendered['status']==status and rendered['values']==[person['store'],person['name'],score,person['case_count'] if person['case_count'] is not None else '待核',needed], 'QIS_PERSONAL_RENDER_VALUE_MISMATCH'
    for entry in receipt['textChecks']:
        l,t,r,b=entry['box'];w,h=entry['text_bounds']
        assert w<=r-l-12 and h<=b-t-entry['vertical_padding'], 'QIS_PERSONAL_TEXT_OVERFLOW'
    return {'schema':personal['schema'],'cutoffDate':personal['cutoffDate'],'sourceSha256':personal['sourceSha256'],
            'rows':len(personal['rows']),'below98':len(personal['below98']),'noData':len(personal['noData']),
            'atLeast98':sum(p['status']=='at_least_98' for p in personal['rows']),'pending':len(personal['pending']),
            'pngSha256':receipt['pngSha256'],'rendererSha256':receipt['rendererSha256']}
