"""Verify source values, missing baselines, formulas and dynamic attachments."""
import hashlib,json,os
from pathlib import Path
from openpyxl import load_workbook
from source_semantics import verify as verify_source_semantics
r=Path(os.environ['RUN_ROOT']);out=Path(os.environ['OUT_DIR']);date=os.environ['REPORT_DATE'];stem=os.environ['SOURCE_STEM']
s=json.loads((r/f'{stem}-source-values.json').read_text());k=json.loads((r/f'kpicalc-{stem}-corrected.json').read_text());d=json.loads((r/'daily-kpi-data.json').read_text());c=json.loads((r/f'closure-audit-{stem}.json').read_text())
b=out/f'TWM_North12B_Daily_Report_{date}.xlsx';cb=out/f'TWM_North12B_Closure_Rate_{date}.xlsx'
v=load_workbook(b,data_only=True);f=load_workbook(b,data_only=False);cv=load_workbook(cb,data_only=True)
assert len(k['items'])==25 and len(k['stores'])==9 and len(k['persons'])>0
insurance=s.get('搭售達成率 (手機保險)',{}).get('values',[]);by={row[0]:row for row in insurance if str(row[0] or '').startswith('DNB')};agg=next((row for row in insurance if row[0]=='北一二B'),None)
for index,(store,row) in enumerate([('北一二B整體',agg)]+[(x['name'],by.get(x['code'])) for x in k['stores']],3):
    for output,sourcecol in [(3,6),(6,7),(8,2),(9,5)]:
        actual=v['手機保險'].cell(index,output).value;expected=row[sourcecol] if row else None
        assert actual==expected if expected is not None else actual=='尚未有資料',f'DISPLAY_VALUE_MISSING insurance {index}/{output}'
    for output,current,prior in [(4,3,2),(7,6,5)]:
        if os.environ['INSURANCE_DOD_SAME_MONTH']!='true' or not isinstance(v['手機保險'].cell(index,prior).value,(int,float)):
            assert v['手機保險'].cell(index,output).value=='暫不比較'
if not d['meta']['sameMonthDod']:
    assert all(v['主力KPI'].cell(row,8).value=='暫不比較' for row in range(3,13))
    assert all(v['加掛得分'].cell(row,7).value=='暫不比較' for row in range(3,13))
assert sum(v[name].max_row-6 for name in ['店長（含代理）','副店長','業代（含銷售人員）'])>=len(k['persons'])
source_contract_version=verify_source_semantics(v,k,s,d,date)
for row in range(4,14):
    for col in range(1,12):assert v['締結率'].cell(row,col).value==cv['締結率近三日'].cell(row,col).value,'CLOSURE_REPORT_MISMATCH'
assert c['latest']['formal_cases']+c['latest']['unassigned_cases']+c['latest']['anomaly_cases']==c['latest']['gap']
sales=d['salesMetrics']
assert len(sales['rows'])==10 and v['配件包膜銷售'].max_row==16 and v['配件包膜銷售'].max_column==4
sales_stores=sorted(sales['rows'][1:],key=lambda row:(not isinstance(row['auto_film_month_count'],(int,float)),-row['auto_film_month_count'] if isinstance(row['auto_film_month_count'],(int,float)) else 0))
sales_fields=['accessory_amount','film_amount','auto_film_month_count']
sales_best={key:max([0]+[row[key] for row in sales_stores if isinstance(row[key],(int,float))]) for key in sales_fields}
for index,row in enumerate([sales['rows'][0]]+sales_stores,4):
    expected=[row['store'],row['accessory_amount'],row['film_amount'],row['auto_film_month_count']]
    for col,value in enumerate(expected,1):
        assert v['配件包膜銷售'].cell(index,col).value==('尚未有資料' if value is None else value),'SALES_REPORT_MISMATCH'
        if index>4 and col>1 and isinstance(value,(int,float)):
            key=sales_fields[col-2]
            cell=f['配件包膜銷售'].cell(index,col)
            fill=cell.fill.fgColor.rgb[-6:];color=cell.font.color.rgb[-6:]
            if value==0:
                assert (fill,color)==('FEE2E2','DC2626'),'SALES_ZERO_COLOR_MISMATCH'
            elif sales_best[key]>0 and value==sales_best[key]:
                assert (fill,color)==('DCFCE7','166534'),'SALES_BEST_COLOR_MISMATCH'
            else:
                assert fill not in ('FEE2E2','DCFCE7'),'SALES_NEUTRAL_COLOR_MISMATCH'
for ws in v:
    for row in ws:
        for cell in row:assert not (cell.data_type=='e'),'FORMULA_ERROR '+ws.title+'!'+cell.coordinate
# Output files follow worksheet presence and the verified paginated receipt.
images={'主力KPI':'Main_KPI','加掛得分':'Addon_Score','手機保險':'Insurance','QIS店績':'QIS','締結率':'Closure_Rate','加減分日目標':'Daily_Targets','店長（含代理）':'Personal_Manager','副店長':'Personal_Deputy','業代（含銷售人員）':'Personal_Sales'}
files=[(b,list(s)),(cb,['續約率(分子)','續約率(分母)'])]+[(out/f'TWM_North12B_{name}_{date}.png',[sheet]) for sheet,name in images.items() if sheet in v]
files.append((out/f'TWM_North12B_Sales_Overview_{date}.png',['上線數KPI_每日上線','AQ其他-包膜與保貼營收']))
receipt=json.loads((r.parent/'private-run/goodspeed-receipt.json').read_text())
assert receipt['status']=='pass'
for item in receipt['files']:
    path=Path(item.get('path') or item.get('file') or item.get('file_name') or '')
    if not path.is_absolute():path=out/path
    files.append((path,['好速案銷售點數']))
entries=[]
for path,sheets in files:
    data=path.read_bytes();assert data
    entries.append({'fileName':path.name,'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest(),'sourceSheets':sheets,'renderStatus':'PENDING_VISUAL_REVIEW' if path.suffix=='.png' else 'WORKBOOK_VALIDATED'})
assert len({x['fileName'] for x in entries})==len(entries)
qis_aggregate=next((row for row in s.get('QIS店績',{}).get('values',[]) if row[0]=='北一二區' and row[2]=='北一二B' and row[3]==9),None)
contract={'status':'pass','version':source_contract_version,'reportDate':date,'sourceSha256':k['meta']['sourceSha256'],
    'dataCutoffDate':k['meta']['data_as_of_date'],'qisRawScore':qis_aggregate[5] if qis_aggregate else None,
    'qisSourceDate':k['meta']['qis_data_as_of_date'],'managerBasis':'personal-employee-kpi' if date>='2026-09-01' else 'legacy',
    'addonRankBasis':'source-company-rank','artifactHashes':{e['fileName']:e['sha256'] for e in entries}}
contract_bytes=json.dumps(contract,ensure_ascii=False,indent=2).encode()
(out/'source-contract-receipt.json').write_bytes(contract_bytes)
(out/'attachment-manifest.json').write_text(json.dumps({'sourceContractVersion':source_contract_version,
    'sourceContractReceiptSha256':hashlib.sha256(contract_bytes).hexdigest(),'reportDate':date,'sourceFile':os.environ['SOURCE_FILE'],
    'sourceSha256':k['meta']['sourceSha256'],'requiredCount':len(entries),'requiredAttachments':entries},ensure_ascii=False,indent=2))
print(json.dumps({'status':'pass','attachments':len(entries),'roster':len(k['persons']),'dodNote':'前日無資料時暫不比較'},ensure_ascii=False))
