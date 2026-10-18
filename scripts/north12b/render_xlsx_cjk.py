#!/usr/bin/env python3
"""Render the approved XLSX sheets with a bundled Traditional-Chinese font.

The workbook is the source of values and styles.  This renderer is deliberately
independent from the spreadsheet renderer that dropped CJK glyphs in staging.
"""

from __future__ import annotations

import math
import hashlib
import json
import os
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(os.environ.get("RUN_ROOT", Path(__file__).resolve().parent))
REPORT_DATE = os.environ["REPORT_DATE"]
OUT = Path(os.environ.get("OUT_DIR", ROOT / "output"))
BOOK = OUT / f"TWM_North12B_Daily_Report_{REPORT_DATE}.xlsx"
FONT_DIR = Path(os.environ.get("FONT_DIR", ROOT / "font-assets"))
FONT_REGULAR = FONT_DIR / "NotoSansTC-Regular.ttf"
FONT_BOLD = FONT_DIR / "NotoSansTC-Bold.ttf"
SCALE = 1.65

SHEETS = {
    "配件包膜銷售": f"TWM_North12B_Sales_Overview_{REPORT_DATE}.png",
    "主力KPI": f"TWM_North12B_Main_KPI_{REPORT_DATE}.png",
    "加掛得分": f"TWM_North12B_Addon_Score_{REPORT_DATE}.png",
    "好速上線明細": f"TWM_North12B_Goodspeed_Detail_{REPORT_DATE}.png",
    "手機保險": f"TWM_North12B_Insurance_{REPORT_DATE}.png",
    "QIS店績": f"TWM_North12B_QIS_{REPORT_DATE}.png",
    "締結率": f"TWM_North12B_Closure_Rate_{REPORT_DATE}.png",
    "加減分日目標": f"TWM_North12B_Daily_Targets_{REPORT_DATE}.png",
    "店長（含代理）": f"TWM_North12B_Personal_Manager_{REPORT_DATE}.png",
    "副店長": f"TWM_North12B_Personal_Deputy_{REPORT_DATE}.png",
    "業代（含銷售人員）": f"TWM_North12B_Personal_Sales_{REPORT_DATE}.png",
}

RED_FILL, RED_TEXT = "FFFEE2E2", "FFDC2626"
GREEN_FILL, GREEN_TEXT = "FFDCFCE7", "FF166534"
BLUE_FILL = "FFDBEAFE"
GRAY_FILL = "FFE5E7EB"


def argb_to_rgb(value, default="FFFFFFFF"):
    value = (value or default).upper()
    if len(value) == 8:
        if value[:2] == "00":
            return tuple(int(default[i:i+2], 16) for i in (2, 4, 6))
        value = value[2:]
    if len(value) != 6:
        value = default[-6:]
    return tuple(int(value[i:i+2], 16) for i in (0, 2, 4))


def color_value(color, default):
    if color is None:
        return default
    if color.type == "rgb" and color.rgb:
        return color.rgb
    if color.type == "indexed" and color.indexed is not None:
        palette = {0: "FF000000", 1: "FFFFFFFF", 2: "FFFF0000", 3: "FF00FF00", 4: "FF0000FF"}
        return palette.get(color.indexed, default)
    return default


def cell_fill(cell):
    value = color_value(cell.fill.fgColor, "FFFFFFFF")
    if cell.fill.fill_type is None or value.startswith("00"):
        return "FFFFFFFF"
    return value


def display_value(value, number_format):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if not isinstance(value, (int, float)):
        return str(value)
    fmt = str(number_format or "General")
    if "pp" in fmt:
        return f"{value:+.1f}pp"
    if "%" in fmt:
        decimals = len(re.search(r"0\.(0+)%", fmt).group(1)) if re.search(r"0\.(0+)%", fmt) else 0
        sign = "+" if "+" in fmt and value > 0 else ""
        return f"{sign}{value * 100:.{decimals}f}%"
    if "pp" in fmt:
        decimals = 1 if "0.0" in fmt else 0
        sign = "+" if value > 0 else ""
        return f"{sign}{value:.{decimals}f}pp"
    decimals = 2 if "0.00" in fmt else (1 if "0.0" in fmt else None)
    if decimals is not None:
        sign = "+" if "+" in fmt and value > 0 else ""
        text = f"{sign}{value:,.{decimals}f}" if "," in fmt else f"{sign}{value:.{decimals}f}"
        if value == 0 and '""' in fmt:
            return ""
        return text
    if float(value).is_integer():
        return f"{int(value):,}" if "," in fmt else str(int(value))
    return f"{value:.4f}".rstrip("0").rstrip(".")


def wrap_text(draw, text, font, width):
    lines = []
    for paragraph in str(text).split("\n"):
        if not paragraph:
            lines.append("")
            continue
        current = ""
        for ch in paragraph:
            trial = current + ch
            if current and draw.textlength(trial, font=font) > width:
                lines.append(current)
                current = ch
            else:
                current = trial
        lines.append(current)
    return lines or [""]


def style_override(sheet, row, col, value, fill, text):
    numeric = value if isinstance(value, (int, float)) else None
    if sheet == "加掛得分" and 3 <= row <= 12 and col in (5, 6) and numeric is not None:
        return (GREEN_FILL, GREEN_TEXT) if numeric >= 15 else (RED_FILL, RED_TEXT)
    if sheet == "手機保險" and 3 <= row <= 12:
        if col in (4, 7) and numeric is not None:
            return (GREEN_FILL, GREEN_TEXT) if numeric >= 0 else (RED_FILL, RED_TEXT)
        if col in (5, 6) and numeric is not None:
            return ("FFF8FAFC", "FF0F172A") if numeric >= 0.5 else (RED_FILL, RED_TEXT)
        if col in (2, 3) and numeric is not None:
            return (GREEN_FILL, GREEN_TEXT) if numeric >= 1 else (RED_FILL, RED_TEXT)
    if sheet == "加減分日目標" and 3 <= row <= 22:
        if row % 2 == 1 and col >= 4:
            if value == "—" or value is None:
                return GRAY_FILL, "FF4B5563"
            return fill, text
        if row % 2 == 0 and col == 3 and numeric is not None:
            return (GREEN_FILL, GREEN_TEXT) if numeric >= 15 else (RED_FILL, RED_TEXT)
        if row % 2 == 0 and col >= 4 and numeric is not None:
            return (GREEN_FILL, GREEN_TEXT) if numeric >= 1 else (RED_FILL, RED_TEXT)
    return fill, text

def conditional_style(ws, cell, value, fill, text):
    """Read numeric cellIs formatting from the actual final workbook."""
    if not isinstance(value, (int, float)):
        return fill, text
    candidates=[]
    for ranges, rules in ws.conditional_formatting._cf_rules.items():
        if cell.coordinate not in ranges.sqref:
            continue
        for rule in rules:
            if rule.type != 'cellIs' or not rule.formula:
                continue
            try: n=float(rule.formula[0])
            except (ValueError, TypeError): continue
            ok={'lessThan':value<n,'lessThanOrEqual':value<=n,'greaterThan':value>n,'greaterThanOrEqual':value>=n,'equal':value==n,'notEqual':value!=n}.get(rule.operator,False)
            if ok: candidates.append(rule)
    # Higher-priority rules own each property; zero alerts precede <100%.
    for rule in sorted(candidates,key=lambda x:x.priority,reverse=True):
        if rule.dxf:
            if rule.dxf.fill:
                fg=color_value(rule.dxf.fill.fgColor,'00000000')
                fill=color_value(rule.dxf.fill.bgColor,fill) if fg.startswith('00') else fg
            if rule.dxf.font and rule.dxf.font.color: text=color_value(rule.dxf.font.color,text)
    return fill,text


def content_bounds(ws, values_ws):
    """Return the last visible row/column with actual content.

    openpyxl keeps formatted-but-empty trailing cells in ``max_row`` and
    ``max_column``.  Rendering those cells caused the Goodspeed image to carry
    several empty rows/columns.  Determine the crop from values/formulas and
    non-empty merged anchors instead of worksheet dimensions.
    """
    last_row = last_col = 1
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            value_cell = values_ws.cell(row, col)
            if cell.value not in (None, "") or value_cell.value not in (None, ""):
                last_row = max(last_row, row)
                last_col = max(last_col, col)
    for rng in ws.merged_cells.ranges:
        anchor = ws.cell(rng.min_row, rng.min_col).value
        value_anchor = values_ws.cell(rng.min_row, rng.min_col).value
        if anchor not in (None, "") or value_anchor not in (None, ""):
            last_row = max(last_row, rng.max_row)
            last_col = max(last_col, rng.max_col)
    return last_row, last_col


def dimensions(ws, last_row, last_col):
    widths = []
    for col in range(1, last_col + 1):
        dim = ws.column_dimensions[get_column_letter(col)]
        w = dim.width if dim.width is not None else 9.0
        widths.append(max(42, round((w * 7 + 7) * SCALE)))
    heights = []
    for row in range(1, last_row + 1):
        dim = ws.row_dimensions[row]
        h = dim.height if dim.height is not None else 24.0
        heights.append(max(24, round(h * 1.333 * SCALE)))
    return widths, heights


def qis_landscape_data(ws):
    """Read the already verified workbook; layout changes never recalculate data."""
    rows=list(ws.iter_rows(values_only=True))
    def at(row,col):return rows[row-1][col-1] if col<=len(rows[row-1]) else None
    store_headers=[i for i,r in enumerate(rows,1) if len(r)>=6 and r[0]=='店點'
        and '達成率' in str(r[1]) and '計件數' in str(r[2]) and '原始分數' in str(r[5])]
    personal_headers=[i for i,r in enumerate(rows,1) if list(r[:5])==['店點','同仁','目前分數','計件數','追加100分筆數']]
    if len(store_headers)!=1 or not personal_headers:raise ValueError('QIS_LANDSCAPE_HEADER_MISMATCH')
    stores=[]
    for row in range(store_headers[0]+1,store_headers[0]+11):
        values=[at(row,col) for col in range(1,7)]
        if not isinstance(values[0],str) or not all(isinstance(v,(int,float)) and not isinstance(v,bool) for v in values[1:]):
            raise ValueError('QIS_LANDSCAPE_STORE_VALUE_MISSING')
        stores.append({'source_row':row,'values':values})
    people=[]
    for header in personal_headers:
        row=header+1
        while row<=len(rows) and isinstance(at(row,2),str) and at(row,2):
            values=[at(row,col) for col in range(1,6)]
            status='no_data' if values[2]=='尚無資料' and type(values[3]) is int and values[3]==0 else 'pending' if '待核' in values else 'below_98'
            if status=='below_98' and not (isinstance(values[2],(int,float)) and not isinstance(values[2],bool) and values[2]<98
                and type(values[3]) is int and values[3]>0 and type(values[4]) is int and values[4]>0):
                raise ValueError('QIS_LANDSCAPE_PERSON_VALUE_MISMATCH')
            people.append({'source_row':row,'values':values,'status':status});row+=1
    if len({p['source_row'] for p in people})!=len(people):raise ValueError('QIS_LANDSCAPE_DUPLICATE_ROW')
    band=next((str(r[0]) for r in rows if str(r[0]).startswith('個人未達98追分')),None)
    match=re.search(r'截至 (\d{4}/\d{2}/\d{2})',band or '')
    personal_cutoff=match[1] if match else None
    store_notes=' '.join(str(r[0]) for r in rows[:personal_headers[0]-2])
    match=re.search(r'截至(\d{4}-\d{2}-\d{2})',store_notes)
    store_cutoff=match[1].replace('-','/') if match else None
    source_match=re.search(r'來源：([^\s]+\.xlsx)',store_notes)
    source_file=source_match[1] if source_match else None
    coverage=next((str(r[0]) for r in rows if str(r[0]).startswith('QIS個績共')),None)
    if not personal_cutoff or not store_cutoff or not coverage or not source_file:raise ValueError('QIS_LANDSCAPE_SOURCE_NOTE_MISSING')
    return {'stores':stores,'people':people,'store_cutoff':store_cutoff,'personal_cutoff':personal_cutoff,'coverage':coverage,'source_file':source_file}


def qis_landscape_layout(data):
    """Two personal columns, complete rows, fixed large fonts, no clipping/shrink."""
    width,height=1920,1080
    left=[166,108,68,100,100,106]
    right=[148,88,114,66,172]
    panels=[{'x':24,'widths':left,'rows':data['stores'],'kind':'stores'}]
    split=(len(data['people'])+1)//2
    for x,people in [(696,data['people'][:split]),(1308,data['people'][split:])]:
        panels.append({'x':x,'widths':right,'rows':people,'kind':'people'})
    for panel in panels:
        count=len(panel['rows'])
        panel['row_height']=60 if panel['kind']=='stores' else 70 if count<=9 else min(70,638//count)
        panel['compact']=panel['kind']=='people' and count>9
        if panel['kind']=='people' and panel['row_height']<38:
            raise ValueError('QIS_LANDSCAPE_ROWS_EXCEED_READABLE_CANVAS')
        if 268+len(panel['rows'])*panel['row_height']>906:
            raise ValueError('QIS_LANDSCAPE_ROWS_EXCEED_READABLE_CANVAS')
    return {'width':width,'height':height,'panels':panels}


def render_qis_landscape(ws, values_ws, output):
    data=qis_landscape_data(values_ws);layout=qis_landscape_layout(data)
    image=Image.new('RGB',(layout['width'],layout['height']),'#F8FAFC');draw=ImageDraw.Draw(image)
    navy,yellow,ink,muted='#17365D','#FFCF17','#0F172A','#475569'
    checked=[];drawn=[]
    def label(box,text,size=26,color=ink,vertical_padding=10):
        # A hard overflow gate preserves the selected font size; never squeeze a row.
        font=ImageFont.truetype(str(FONT_BOLD),size)
        bounds=draw.multiline_textbbox((0,0),str(text),font=font,spacing=3,align='center')
        w,h=bounds[2]-bounds[0],bounds[3]-bounds[1]
        if w>box[2]-box[0]-12 or h>box[3]-box[1]-vertical_padding:
            raise ValueError('QIS_LANDSCAPE_TEXT_OVERFLOW '+str(text))
        x=(box[0]+box[2]-w)/2-bounds[0];y=(box[1]+box[3]-h)/2-bounds[1]
        draw.multiline_text((x,y),str(text),font=font,fill=color,spacing=3,align='center')
        checked.append({'text':str(text),'box':box,'font_px':size,'text_bounds':[w,h],'vertical_padding':vertical_padding})
    draw.rounded_rectangle((24,20,1896,106),radius=12,fill=navy)
    label((40,20,1130,106),'北一二B QIS 戰報',44,'#FFFFFF')
    label((1320,20,1880,106),REPORT_DATE.replace('-','/'),36,'#FFFFFF')
    draw.rectangle((24,128,672,192),fill=navy)
    label((24,128,672,192),'店績｜截至 '+data['store_cutoff'][5:],30,'#FFFFFF')
    draw.rectangle((696,128,1896,192),fill=navy)
    low=sum(p['status']=='below_98' for p in data['people']);empty=sum(p['status']=='no_data' for p in data['people'])
    label((704,128,1310,192),'個人追98｜截至 '+data['personal_cutoff'][5:],30,'#FFFFFF')
    label((1310,128,1888,192),f'未達98 {low}人／尚無計件 {empty}人',26,'#FFFFFF')
    for panel in layout['panels']:
        x=panel['x'];widths=panel['widths'];xs=[x]
        for w in widths:xs.append(xs[-1]+w)
        headers=['店點','達成率','計件\n數','服務\n態度','專業\n能力','原始\n分數'] if panel['kind']=='stores' else ['店點','同仁','目前\n分數','計件\n數','追加100分\n筆數']
        for col,heading in enumerate(headers):
            box=(xs[col],192,xs[col+1],268);draw.rectangle(box,fill=yellow,outline='#CBD5E1',width=2);label(box,heading,26,navy)
        for index,row in enumerate(panel['rows']):
            y=268+index*panel['row_height'];values=row['values'];texts=[]
            for col,value in enumerate(values):
                box=(xs[col],y,xs[col+1],y+panel['row_height'])
                fill='#FFF4C2' if panel['kind']=='stores' and index==0 else '#FFFFFF' if index%2 else '#F1F5F9'
                color=ink;size=26
                if panel['kind']=='stores':
                    text=str(value) if col==0 else f'{value*100:.1f}%' if col==1 else str(value) if col==2 else f'{value:.2f}'
                    if col==1:color='#059669' if value>=1 else '#DC2626'
                    if col==5:color='#059669' if value>=98 else '#DC2626';fill='#DCFCE7' if value>=98 else '#FEE2E2'
                else:
                    text=str(value);fill='#FFF1F2' if row['status']=='below_98' else '#E2E8F0'
                    if isinstance(value,float):text=f'{value:.5f}'.rstrip('0').rstrip('.')
                    if col==4:color='#B91C1C' if row['status']=='below_98' else muted;size=38 if isinstance(value,int) else 26
                    if row['status']=='no_data' and col==2:text='無資料' if panel['compact'] else '尚無\n資料';size=26 if panel['compact'] else 24
                    if row['status']=='no_data' and col==4:text='1' if panel['compact'] else '首筆100\n即100';size=38 if panel['compact'] else 26
                draw.rectangle(box,fill=fill,outline='#CBD5E1',width=2);label(box,text,size,color,8 if panel['compact'] else 10);texts.append(text)
            drawn.append({'kind':panel['kind'],'status':row.get('status'),'source_row':row['source_row'],'values':values,'texts':texts})
    notes=[
        '店績整體：達成率、服務態度與專業能力按9店算術平均；計件數加總；原始分數取來源F13。',
        '追98追加100分筆數＝向上取整〔計件數 × (98－目前分數) ÷ 2〕，採等權平均試算。',
        '依報表分數試算，臨界可能差1筆；實際以公司更新為準，未反推原始總分。',
        '來源'+data['source_file']+'；'+data['coverage'],
    ]
    if any(panel['compact'] for panel in layout['panels']) and empty:
        notes[-1]+=' 尚無計件者：首筆100即100。'
    for index,note in enumerate(notes):label((24,906+index*40,1896,946+index*40),note,23,muted)
    image.save(output,format='PNG',compress_level=6)
    return {'status':'pass','dimensions':[image.width,image.height],'aspectRatio':'16:9','singlePng':True,
        'storeRows':len(data['stores']),'personalRows':len(data['people']),'below98':low,'noData':empty,
        'storeCutoff':data['store_cutoff'],'personalCutoff':data['personal_cutoff'],
        'allTextFitsAtSelectedFontSize':True,'fontShrinking':False,'drawnRows':drawn,'textChecks':checked,
        'pngSha256':hashlib.sha256(Path(output).read_bytes()).hexdigest()}


def render_sheet(ws, values_ws, output):
    last_row, last_col = content_bounds(ws, values_ws)
    widths, heights = dimensions(ws, last_row, last_col)
    xs = [18]
    ys = [18]
    for w in widths:
        xs.append(xs[-1] + w)
    for h in heights:
        ys.append(ys[-1] + h)
    image = Image.new("RGB", (xs[-1] + 18, ys[-1] + 18), (248, 250, 252))
    draw = ImageDraw.Draw(image)
    merged = {}
    children = set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    children.add((r, c))

    for row in range(1, last_row + 1):
        for col in range(1, last_col + 1):
            if (row, col) in children:
                continue
            cell = ws.cell(row, col)
            if isinstance(cell, MergedCell):
                continue
            end_row, end_col = merged.get((row, col), (row, col))
            box = (xs[col - 1], ys[row - 1], xs[end_col], ys[end_row])
            value = values_ws.cell(row, col).value
            if value is None and not (isinstance(cell.value, str) and cell.value.startswith("=")):
                value = cell.value
            fill = cell_fill(cell)
            text_color = color_value(cell.font.color, "FF0F172A")
            fill, text_color = conditional_style(ws, cell, value, fill, text_color)
            fill, text_color = style_override(ws.title, row, col, value, fill, text_color)
            draw.rectangle(box, fill=argb_to_rgb(fill), outline=(203, 213, 225), width=max(1, round(SCALE)))
            text = display_value(value, cell.number_format)
            if not text:
                continue
            size_pt = float(cell.font.sz or 10)
            if row == 1:
                size_pt = max(size_pt, 15)
            px = max(13, round(size_pt * 1.45 * SCALE))
            font_path = FONT_BOLD if cell.font.bold or row <= 3 else FONT_REGULAR
            font = ImageFont.truetype(str(font_path), px)
            padding = round(5 * SCALE)
            available = max(10, box[2] - box[0] - padding * 2)
            if cell.alignment.wrap_text or "\n" in text:
                lines = wrap_text(draw, text, font, available)
            else:
                while px > 11 and draw.textlength(text, font=font) > available:
                    px -= 1
                    font = ImageFont.truetype(str(font_path), px)
                lines = [text]
            line_height = round(px * 1.25)
            block_height = line_height * len(lines)
            while px > 11 and block_height > box[3]-box[1]-padding*2:
                px -= 1
                font = ImageFont.truetype(str(font_path), px)
                lines = wrap_text(draw,text,font,available) if cell.alignment.wrap_text or '\n' in text else [text]
                line_height=round(px*1.25)
                block_height=line_height*len(lines)
            if block_height > box[3]-box[1]-padding*2:
                raise ValueError('RENDER_TEXT_OVERFLOW '+ws.title+' '+cell.coordinate)
            valign = cell.alignment.vertical or "center"
            if valign == "top":
                y = box[1] + padding
            elif valign == "bottom":
                y = box[3] - block_height - padding
            else:
                y = box[1] + max(padding, (box[3] - box[1] - block_height) // 2)
            align = cell.alignment.horizontal or "center"
            for line in lines:
                line_width = draw.textlength(line, font=font)
                if align == "left":
                    x = box[0] + padding
                elif align == "right":
                    x = box[2] - padding - line_width
                else:
                    x = box[0] + (box[2] - box[0] - line_width) / 2
                draw.text((x, y), line, font=font, fill=argb_to_rgb(text_color, "FF0F172A"))
                y += line_height
    image.save(output, format="PNG", compress_level=6)
    print(f"{ws.title}\t{image.width}x{image.height}\t{output.name}")


def main():
    for path in (BOOK, FONT_REGULAR, FONT_BOLD):
        if not path.exists():
            raise SystemExit(f"missing: {path}")
    wb = load_workbook(BOOK, data_only=False, keep_links=False)
    values = load_workbook(BOOK, data_only=True, keep_links=False)
    for sheet, filename in SHEETS.items():
        if sheet == "好速上線明細": continue
        if sheet == 'QIS店績':
            receipt=render_qis_landscape(wb[sheet],values[sheet],OUT/filename)
            receipt['workbookSha256']=hashlib.sha256(BOOK.read_bytes()).hexdigest()
            receipt['rendererSha256']=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
            receipt['fontSha256']={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [FONT_REGULAR,FONT_BOLD]}
            private=ROOT.parent/'private-run';private.mkdir(parents=True,exist_ok=True)
            (private/'qis-landscape-render.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2))
            print(f"{sheet}\t1920x1080\t{filename}")
            continue
        render_sheet(wb[sheet], values[sheet], OUT / filename)
    wb.close()
    values.close()


if __name__ == "__main__":
    main()
