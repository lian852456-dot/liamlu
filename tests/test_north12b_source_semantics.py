import importlib.util, unittest
from pathlib import Path
from openpyxl import Workbook
spec=importlib.util.spec_from_file_location('source_semantics',Path(__file__).resolve().parents[1]/'scripts/north12b/source_semantics.py')
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
class SourceSemanticsTest(unittest.TestCase):
    def setUp(self):
        self.b=Workbook();self.b.remove(self.b.active)
        for name in ['店長（含代理）','加掛得分','主力KPI','QIS店績']:self.b.create_sheet(name)
        item={'a':10,'reportRate':1.25}
        self.p={'role':'店長','official':2.1846,'items':{k:dict(item) for k in m.PERSONAL_METRICS}}
        self.k={'persons':[self.p],'aggregateRates':{'預付卡開卡面額':1.15},'stores':[{'name':'合成店','items':{'預付卡開卡面額':{'reportRate':1.162}}}]}
        self.d={'summary':[{'store':'北一二B整體','rank':10},{'store':'合成店','rank':290}]}
        self.b['店長（含代理）'].cell(5,4,2.1846)
        for i,_ in enumerate(m.PERSONAL_METRICS):self.b['店長（含代理）'].cell(5,6+i*3,10);self.b['店長（含代理）'].cell(5,7+i*3,1.25)
        self.b['加掛得分'].cell(2,16,'Apple/Google服務及雜誌週刊開通數')
        for rn,rank,prepaid in [(3,10,1.15),(4,290,1.162)]:
            self.b['加掛得分'].cell(rn,3,rank);self.b['主力KPI'].cell(rn,16,prepaid)
        self.b['QIS店績'].cell(3,6,97.21333)
        self.s={'QIS店績':{'values':[['北一二區',None,'北一二B',9,None,97.21333]]}}
    def check(self):return m.verify(self.b,self.k,self.s,self.d,'2026-10-13')
    def test_correct_source_semantics(self):self.assertEqual(self.check(),m.VERSION)
    def test_manager_store_override_rejected(self):
        self.b['店長（含代理）'].cell(5,4,1.057)
        with self.assertRaisesRegex(AssertionError,'MANAGER_PERSONAL_OVERALL'):self.check()
    def test_prepaid_rt_and_aggregate_header_confusion_rejected(self):
        self.b['主力KPI'].cell(4,16,0.87)
        with self.assertRaisesRegex(AssertionError,'MAIN_PREPAID'):self.check()
        self.b['主力KPI'].cell(4,16,1.162);self.b['主力KPI'].cell(3,16,1.162)
        with self.assertRaisesRegex(AssertionError,'MAIN_PREPAID'):self.check()
    def test_qis_weighted_mean_rejected(self):
        self.b['QIS店績'].cell(3,6,97.5)
        with self.assertRaisesRegex(AssertionError,'QIS_SOURCE'):self.check()
    def test_ordinal_rank_and_product_alias_rejected(self):
        self.b['加掛得分'].cell(4,3,1)
        with self.assertRaisesRegex(AssertionError,'ADDON_COMPANY'):self.check()
        self.b['加掛得分'].cell(4,3,290);self.b['加掛得分'].cell(2,16,'YT/天下')
        with self.assertRaisesRegex(AssertionError,'ADDON_SERVICE'):self.check()
if __name__=='__main__':unittest.main()
