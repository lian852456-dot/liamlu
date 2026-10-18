import hashlib
import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


MODULE_PATH = Path(__file__).resolve().parents[1] / "scripts/north12b/qis_personal.py"
SPEC = importlib.util.spec_from_file_location("qis_personal", MODULE_PATH)
qis = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(qis)


class PersonalQisTest(unittest.TestCase):
    def test_reported_average_scenarios_use_decimal_ceiling(self):
        cases = [
            (10, "97.8", 1),
            (5, "95.6", 6),
            (3, "96.5", 3),
        ]
        for count, score, expected in cases:
            with self.subTest(count=count, score=score):
                result = qis.estimate_100_cases(count, score)
                self.assertEqual(result["status"], "below_98")
                self.assertEqual(result["additional_100"], expected)
                self.assertGreaterEqual(result["projected_report_average"], 98)
                if expected:
                    previous = (count * float(score) + 100 * (expected - 1)) / (count + expected - 1)
                    self.assertLess(previous, 98)
                self.assertEqual(result["current_score"], float(score))

    def test_at_or_above_target_needs_no_cases(self):
        result = qis.estimate_100_cases(8, "98.00")
        self.assertEqual(result["status"], "at_least_98")
        self.assertEqual(result["additional_100"], 0)
        self.assertEqual(result["projected_report_average"], 98.0)

    def test_zero_count_is_no_data_and_one_case_scenario(self):
        for score in (None, 0):
            with self.subTest(score=score):
                result = qis.estimate_100_cases(0, score)
                self.assertEqual(result["status"], "no_data")
                self.assertIsNone(result["current_score"])
                self.assertEqual(result["case_count"], 0)
                self.assertEqual(result["additional_100"], 1)
                self.assertEqual(result["projected_report_average"], 100.0)

    def test_missing_or_invalid_inputs_are_pending_not_zero(self):
        for count, score in ((5, None), (None, 97), ("bad", 97), (-1, 80), (1.5, 80), (5, -1), (5, 100.1)):
            with self.subTest(count=count, score=score):
                result = qis.estimate_100_cases(count, score)
                self.assertEqual(result["status"], "pending")
                self.assertIsNone(result["additional_100"])
                self.assertIsNone(result["projected_report_average"])
        self.assertIsNone(qis.estimate_100_cases(5, None)["current_score"])

    def roster(self):
        return {
            "persons": [
                {"employee_id": "E001", "pname": "合成甲", "store": "S1", "store_name": "合成酒泉"},
                {"employee_id": "E002", "pname": "合成乙", "store": "S2", "store_name": "合成復興"},
                {"employee_id": "E003", "pname": "合成丙", "store": "S1", "store_name": "合成酒泉"},
            ],
            "stores": [
                {"code": "S1", "name": "合成酒泉"},
                {"code": "S2", "name": "合成復興"},
            ],
        }

    def make_source(self, *, reordered=False, duplicate_header=False, missing_score=False, formula_score=False):
        book = Workbook()
        sheet = book.active
        sheet.title = "QIS個績"
        sheet.cell(1, 1, "2026/10/1~2026/10/13")
        if reordered:
            cols = {"scope": 2, "store": 4, "name": 7, "employee": 9, "count": 11, "score": 5}
        else:
            cols = {"store": 2, "name": 3, "employee": 4, "scope": 5, "count": 8, "score": 13}
        header = 6
        for key, label in (("store", "店點"), ("name", "姓名"), ("employee", "員編"),
                           ("scope", "區域"), ("count", "計件數"), ("score", "原始分數")):
            if key == "score" and missing_score:
                continue
            sheet.cell(header, cols[key], label)
        if duplicate_header:
            sheet.cell(header, cols["score"] + 1, "原始分數")
        # These raw-score subheaders must not be mistaken for another main header.
        sheet.cell(7, cols["score"] - 1, "原始分數")
        sheet.cell(7, cols["score"] - 2, "原始分數")
        data = [
            ("S1", "合成甲", "E001", "北一二B", 10, 97.8),
            ("S2", "合成乙", "E002", "北一二B", 0, 0),
            ("S1", "合成丙", "E003", "北一二B", 5, None),
            ("S1", "外區人員", "OUT1", "北一二A", 4, 70),
        ]
        for rn, item in enumerate(data, 8):
            mapping = dict(zip(("store", "name", "employee", "scope", "count", "score"), item))
            for key, value in mapping.items():
                if key == "score" and value is None:
                    continue
                if key == "score" and formula_score and rn == 8:
                    sheet.cell(rn, cols[key], "=1+2")
                else:
                    sheet.cell(rn, cols[key], value)
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.close()
        book.save(tmp.name)
        return Path(tmp.name), cols

    def parse(self, **kwargs):
        path, _ = self.make_source(**kwargs)
        self.addCleanup(path.unlink, missing_ok=True)
        return path, qis.extract_personal_qis(path, self.roster(), "2026-10-13")

    def test_extracts_scoped_rows_and_discovers_reordered_identity_and_metric_columns(self):
        for reordered in (False, True):
            with self.subTest(reordered=reordered):
                path, report = self.parse(reordered=reordered)
                self.assertEqual(report["schema"], "north12b-personal-qis/v1")
                self.assertEqual(report["cutoffDate"], "2026-10-13")
                self.assertEqual(report["sourceSha256"], hashlib.sha256(path.read_bytes()).hexdigest())
                self.assertEqual(report["header"]["row"], 6)
                self.assertEqual(len(report["rows"]), 3)
                self.assertEqual([row["employee_id"] for row in report["rows"]], ["E001", "E002", "E003"])
                self.assertEqual(report["rows"][0]["case_count"], 10)
                self.assertEqual(report["rows"][0]["current_score"], 97.8)
                self.assertEqual(report["rows"][0]["score_cell"], "M8" if not reordered else "E8")
                self.assertEqual(report["rows"][1]["status"], "no_data")
                self.assertEqual(report["rows"][2]["status"], "pending")
                self.assertEqual(len(report["below98"]), 1)
                self.assertEqual(len(report["noData"]), 1)
                self.assertEqual(len(report["pending"]), 1)
                self.assertEqual(report["rows"][0]["store"], "合成酒泉")

    def test_formula_without_cached_source_value_stays_pending(self):
        _, report = self.parse(formula_score=True)
        row = report["rows"][0]
        self.assertEqual(row["status"], "pending")
        self.assertEqual(row["score_value_kind"], "formula-cached")
        self.assertIsNone(row["current_score"])
        self.assertIsNone(row["additional_100"])

    def test_boolean_source_count_is_pending_not_one_case(self):
        path, cols = self.make_source()
        self.addCleanup(path.unlink, missing_ok=True)
        book = load_workbook(path)
        book['QIS個績'].cell(8, cols['count'], True)
        book.save(path)
        report = qis.extract_personal_qis(path, self.roster(), '2026-10-13')
        self.assertEqual(report['rows'][0]['status'], 'pending')
        self.assertIsNone(report['rows'][0]['case_count'])
        self.assertIsNone(report['rows'][0]['additional_100'])

    def test_missing_or_multiple_main_metric_headers_fail_closed(self):
        for options in ({"missing_score": True}, {"duplicate_header": True}):
            path, _ = self.make_source(**options)
            self.addCleanup(path.unlink, missing_ok=True)
            with self.subTest(options=options), self.assertRaisesRegex(ValueError, "QIS_PERSONAL_HEADER_MISMATCH"):
                qis.extract_personal_qis(path, self.roster(), "2026-10-13")

    def test_duplicate_employee_in_scope_fails_closed(self):
        path, _ = self.make_source()
        self.addCleanup(path.unlink, missing_ok=True)
        book = Workbook()
        sheet = book.active
        sheet.title = "QIS個績"
        # Recreate a minimal source with the same employee appearing twice.
        sheet.cell(1, 1, "2026/10/1~2026/10/13")
        for col, value in ((2, "店點"), (3, "姓名"), (4, "員編"), (5, "區域"), (8, "計件數"), (13, "原始分數")):
            sheet.cell(6, col, value)
        for rn in (8, 9):
            for col, value in ((2, "S1"), (3, "合成甲"), (4, "E001"), (5, "北一二B"), (8, 10), (13, 97)):
                sheet.cell(rn, col, value)
        book.save(path)
        with self.assertRaisesRegex(ValueError, "QIS_PERSONAL_DUPLICATE_EMPLOYEE"):
            qis.extract_personal_qis(path, self.roster(), "2026-10-13")


if __name__ == "__main__":
    unittest.main()
