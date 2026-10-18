import importlib.util
import os
import unittest
from pathlib import Path
from openpyxl import Workbook

os.environ.setdefault('REPORT_DATE','2026-10-13')
path=Path(__file__).resolve().parents[1]/'scripts/north12b/render_xlsx_cjk.py'
spec=importlib.util.spec_from_file_location('qis_landscape_renderer',path)
renderer=importlib.util.module_from_spec(spec);spec.loader.exec_module(renderer)


def fixture():
    sheet=Workbook().active
    sheet['A1']='北一二B QIS店績｜2026-10-13'
    for col,label in enumerate(['店點','達成率','計件數','服務態度','專業能力','原始分數'],1):sheet.cell(2,col,label)
    for row in range(3,13):
        for col,value in enumerate([f'合成店{row}',1.5,12,97.72,98.01,97.865],1):sheet.cell(row,col,value)
    sheet['A14']='來源：1013.xlsx QIS店績，截至2026-10-10；整體按9店平均。'
    sheet['A16']='個人未達98追分｜截至 2026/10/09｜17人'
    labels=['店點','同仁','目前分數','計件數','追加100分筆數']
    for col,label in enumerate(labels,1):sheet.cell(17,col,label)
    for row in range(18,35):
        for col,value in enumerate(['合成店',f'合成{row}',97.72,22,4],1):sheet.cell(row,col,value)
    sheet['A36']='尚無計件｜1人'
    for col,label in enumerate(labels,1):sheet.cell(37,col,label)
    for col,value in enumerate(['合成店','無計件同仁','尚無資料',0,'首筆100即100'],1):sheet.cell(38,col,value)
    sheet['A40']='QIS個績共32人：未達98 17人、尚無計件 1人、待核 0人。'
    return sheet


class QisLandscapeTests(unittest.TestCase):
    def test_exact_ten_store_rows_and_eighteen_people_keep_values_and_dates(self):
        data=renderer.qis_landscape_data(fixture())
        self.assertEqual(len(data['stores']),10)
        self.assertEqual(len(data['people']),18)
        self.assertEqual(data['stores'][0]['values'],['合成店3',1.5,12,97.72,98.01,97.865])
        self.assertEqual(data['people'][0]['values'],['合成店','合成18',97.72,22,4])
        self.assertEqual(data['people'][-1]['values'],['合成店','無計件同仁','尚無資料',0,'首筆100即100'])
        self.assertEqual(data['people'][-1]['status'],'no_data')
        self.assertEqual((data['store_cutoff'],data['personal_cutoff']),('2026/10/10','2026/10/09'))

    def test_two_personal_columns_include_each_source_row_once_in_16_by_9(self):
        data=renderer.qis_landscape_data(fixture());layout=renderer.qis_landscape_layout(data)
        self.assertEqual((layout['width'],layout['height']),(1920,1080))
        panels=layout['panels'];self.assertEqual([len(p['rows']) for p in panels],[10,9,9])
        actual=[p['source_row'] for panel in panels[1:] for p in panel['rows']]
        self.assertEqual(actual,[p['source_row'] for p in data['people']])
        for panel in panels:
            self.assertGreaterEqual(panel['x'],24)
            self.assertLessEqual(panel['x']+sum(panel['widths']),1896)
            self.assertLessEqual(268+len(panel['rows'])*panel['row_height'],906)
        for left,right in zip(panels,panels[1:]):
            self.assertLess(left['x']+sum(left['widths']),right['x'])

    def test_overflow_stops_instead_of_dropping_people_or_shrinking(self):
        data=renderer.qis_landscape_data(fixture())
        data['people']*=2
        with self.assertRaisesRegex(ValueError,'ROWS_EXCEED_READABLE_CANVAS'):
            renderer.qis_landscape_layout(data)

    def test_entire_32_person_population_fits_without_changing_font_sizes(self):
        data=renderer.qis_landscape_data(fixture())
        data['people']=[{**data['people'][0],'source_row':n} for n in range(31)]+[data['people'][-1]]
        layout=renderer.qis_landscape_layout(data)
        self.assertEqual([len(p['rows']) for p in layout['panels']],[10,16,16])
        self.assertEqual(sum(len(p['rows']) for p in layout['panels'][1:]),32)
        for panel in layout['panels'][1:]:
            self.assertGreaterEqual(panel['row_height'],38)
            self.assertLessEqual(268+len(panel['rows'])*panel['row_height'],906)

    def test_missing_cutoff_and_unexpected_person_values_stop(self):
        sheet=fixture();sheet['A14']='no source date'
        with self.assertRaisesRegex(ValueError,'SOURCE_NOTE_MISSING'):renderer.qis_landscape_data(sheet)
        sheet=fixture();sheet['E18']=None
        with self.assertRaisesRegex(ValueError,'PERSON_VALUE_MISMATCH'):renderer.qis_landscape_data(sheet)


if __name__=='__main__':unittest.main()
